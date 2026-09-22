"""Attendance export.

The old route wrote a timestamped file into ``exports/`` on every request, served it
from disk and never deleted it, and ran an unbounded query.  Exports are now streamed
from memory and bounded by ``EXPORT_MAX_ROWS``, so the things worth pinning are: the
bytes are a well-formed CSV/XLSX, the response is an attachment, the row cap is
honoured, and an empty range does not hand back an empty file as if it were data.
"""

from __future__ import annotations

import csv
import io

import pytest
from openpyxl import load_workbook

from app.services import export_service
from tests.conftest import app_today


def parse_csv(payload: bytes) -> list[list[str]]:
    text = payload.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


class TestCsv:
    def test_header_matches_the_declared_columns(self, app, attendance_history):  # noqa: ARG002
        rows = parse_csv(export_service.to_csv_bytes(attendance_history))
        assert rows[0] == list(export_service.HEADERS)

    def test_one_row_per_record(self, app, attendance_history):  # noqa: ARG002
        rows = parse_csv(export_service.to_csv_bytes(attendance_history))
        assert len(rows) == len(attendance_history) + 1

    def test_starts_with_a_bom_so_excel_reads_utf8(self, app, attendance_history):  # noqa: ARG002
        """Without the BOM Excel mis-decodes non-ASCII names on a double-click open."""
        assert export_service.to_csv_bytes(attendance_history).startswith(b"\xef\xbb\xbf")

    def test_student_columns_are_resolved(self, app, students, attendance_history):  # noqa: ARG002
        rows = parse_csv(export_service.to_csv_bytes(attendance_history))
        roll_column = export_service.HEADERS.index("Student ID")
        assert {row[roll_column] for row in rows[1:]} == {"S000", "S001", "S002"}

    def test_marked_by_falls_back_to_system(self, app, students, db):
        from app.models import AttendanceRecord

        record = AttendanceRecord(student_id=students[0].id, date=app_today(), status="Present")
        db.session.add(record)
        db.session.commit()

        rows = parse_csv(export_service.to_csv_bytes([record]))
        assert rows[1][export_service.HEADERS.index("Marked By")] == "System"

    def test_an_empty_record_list_still_has_a_header(self, app):  # noqa: ARG002
        assert parse_csv(export_service.to_csv_bytes([])) == [list(export_service.HEADERS)]


class TestXlsx:
    def test_it_is_a_readable_workbook(self, app, attendance_history):  # noqa: ARG002
        book = load_workbook(io.BytesIO(export_service.to_xlsx_bytes(attendance_history)))
        assert book.active.title == "Attendance Records"

    def test_header_row_is_frozen(self, app, attendance_history):  # noqa: ARG002
        book = load_workbook(io.BytesIO(export_service.to_xlsx_bytes(attendance_history)))
        assert book.active.freeze_panes == "A2"

    def test_every_record_gets_a_row(self, app, attendance_history):  # noqa: ARG002
        sheet = load_workbook(io.BytesIO(export_service.to_xlsx_bytes(attendance_history))).active
        values = [
            sheet.cell(row=1, column=col).value for col in range(1, len(export_service.HEADERS) + 1)
        ]
        assert values == list(export_service.HEADERS)
        # Header + one row per record, then a blank line before the summary block.
        assert sheet.cell(row=len(attendance_history) + 1, column=1).value is not None

    def test_summary_block_totals_the_records(self, app, attendance_history):  # noqa: ARG002
        sheet = load_workbook(io.BytesIO(export_service.to_xlsx_bytes(attendance_history))).active
        summary_row = len(attendance_history) + 3
        assert sheet.cell(row=summary_row, column=1).value == "Summary"
        assert sheet.cell(row=summary_row + 1, column=1).value == (
            f"Total records: {len(attendance_history)}"
        )

    def test_an_empty_workbook_is_still_valid(self, app):  # noqa: ARG002
        book = load_workbook(io.BytesIO(export_service.to_xlsx_bytes([])))
        assert book.active.cell(row=1, column=1).value == "Date"


class TestBuildExport:
    @pytest.mark.parametrize(
        ("fmt", "extension", "mimetype_fragment"),
        [
            ("csv", "csv", "text/csv"),
            ("excel", "xlsx", "spreadsheetml"),
            # Anything unrecognised falls back to CSV rather than erroring.
            ("pdf", "csv", "text/csv"),
        ],
    )
    def test_format_selection(self, app, attendance_history, fmt, extension, mimetype_fragment):  # noqa: ARG002
        payload, mimetype, ext = export_service.build_export(attendance_history, fmt)
        assert ext == extension
        assert mimetype_fragment in mimetype
        assert payload


class TestRoute:
    def test_csv_download_is_an_attachment(self, admin_client, attendance_history):  # noqa: ARG002
        response = admin_client.get("/export_attendance?format=csv")

        assert response.status_code == 200
        assert "text/csv" in response.headers["Content-Type"]
        assert response.headers["Content-Disposition"].startswith("attachment;")
        assert response.headers["Content-Disposition"].endswith('.csv"')

    def test_xlsx_download_is_an_attachment(self, admin_client, attendance_history):  # noqa: ARG002
        response = admin_client.get("/export_attendance?format=excel")

        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["Content-Type"]
        assert response.headers["Content-Disposition"].endswith('.xlsx"')

    def test_the_payload_round_trips(self, admin_client, attendance_history):  # noqa: ARG002
        response = admin_client.get("/export_attendance?format=csv")
        rows = parse_csv(response.data)
        assert rows[0] == list(export_service.HEADERS)
        assert len(rows) - 1 == int(response.headers["X-Export-Rows"])

    def test_row_count_is_reported_in_a_header(self, admin_client, attendance_history):
        response = admin_client.get("/export_attendance?format=csv")
        assert int(response.headers["X-Export-Rows"]) == len(attendance_history)

    def test_an_empty_range_redirects_instead_of_returning_an_empty_file(
        self, admin_client, attendance_history
    ):  # noqa: ARG002
        response = admin_client.get("/export_attendance?date_from=1999-01-01&date_to=1999-01-02")
        assert response.status_code == 302

    def test_the_row_cap_is_enforced(self, app, admin_client, attendance_history):
        app.config["EXPORT_MAX_ROWS"] = 3
        response = admin_client.get("/export_attendance?format=csv")

        assert response.status_code == 200
        assert int(response.headers["X-Export-Rows"]) == 3
        assert len(parse_csv(response.data)) == 4  # header + 3

    def test_a_bad_date_does_not_500(self, admin_client, attendance_history):  # noqa: ARG002
        response = admin_client.get("/export_attendance?date_from=not-a-date")
        assert response.status_code == 200
